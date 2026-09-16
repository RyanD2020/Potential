"""Utilities for extracting content from uploaded supporting documents.

Two kinds of "supporting document" are handled differently:
- Text-bearing files (PDF, DOCX, TXT, MD, XLSX, CSV, PPTX) get their text
  extracted and folded into one context string sent to the model.
- Image files (PNG, JPG, etc.) are NOT run through OCR - they're kept as
  raw bytes and handed to the model as actual vision input, since OCR would
  lose real information a flowchart or screenshot conveys (boxes, arrows,
  layout) that plain text extraction can't capture.

Keeps document parsing separate from the Streamlit UI and the provider
calls, so each file has one clear job.
"""
from __future__ import annotations

import base64
import io
from typing import Any, Dict, List

MAX_CHARS_PER_DOC = 8000  # keeps prompts a reasonable size even with several docs
MAX_IMAGE_BYTES = 5 * 1024 * 1024  # 5 MB per image, generous for a screenshot/diagram

IMAGE_EXTENSIONS = (".png", ".jpg", ".jpeg", ".gif", ".webp")

_IMAGE_MIME_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
}


def is_image_file(filename: str) -> bool:
    return filename.lower().endswith(IMAGE_EXTENSIONS)


def extract_text_from_file(uploaded_file) -> str:
    """Extract plain text from a Streamlit UploadedFile.

    Supports .txt, .md, .pdf, .docx, .xlsx/.xlsm, .csv, and .pptx. Returns a
    readable error message (instead of raising) for anything else, so a bad
    upload never crashes the app.
    """
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    try:
        if name.endswith(".txt") or name.endswith(".md"):
            return data.decode("utf-8", errors="ignore")

        if name.endswith(".csv"):
            return data.decode("utf-8", errors="ignore")

        if name.endswith(".pdf"):
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            pages = [page.extract_text() or "" for page in reader.pages]
            return "\n\n".join(pages)

        if name.endswith(".docx"):
            import docx

            document = docx.Document(io.BytesIO(data))
            paragraphs = [p.text for p in document.paragraphs]
            return "\n".join(paragraphs)

        if name.endswith(".xlsx") or name.endswith(".xlsm"):
            import openpyxl

            workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
            sheets_text = []
            for sheet in workbook.worksheets:
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    if any(cell is not None for cell in row):
                        rows.append(
                            "\t".join("" if cell is None else str(cell) for cell in row)
                        )
                sheets_text.append(f"[Sheet: {sheet.title}]\n" + "\n".join(rows))
            return "\n\n".join(sheets_text)

        if name.endswith(".xls"):
            return (
                f"[Legacy .xls format for '{uploaded_file.name}' isn't supported - "
                "please re-save it as .xlsx and re-upload.]"
            )

        if name.endswith(".pptx"):
            from pptx import Presentation

            presentation = Presentation(io.BytesIO(data))
            slides_text = []
            for i, slide in enumerate(presentation.slides, start=1):
                lines = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        text = shape.text_frame.text.strip()
                        if text:
                            lines.append(text)
                slides_text.append(f"[Slide {i}]\n" + "\n".join(lines))
            return "\n\n".join(slides_text)

        return (
            f"[Unsupported file type for '{uploaded_file.name}' - "
            "only .txt, .md, .pdf, .docx, .xlsx, .csv, .pptx are read as text.]"
        )

    except Exception as exc:  # noqa: BLE001 - surface any parser error to the user
        return f"[Could not read '{uploaded_file.name}': {exc}]"


def build_supporting_docs_context(uploaded_files: List) -> str:
    """Turn a list of non-image uploaded files into one context string.

    Each document is truncated to MAX_CHARS_PER_DOC so a handful of large
    uploads don't blow out the prompt.
    """
    if not uploaded_files:
        return ""

    sections = []
    for f in uploaded_files:
        text = extract_text_from_file(f)
        if len(text) > MAX_CHARS_PER_DOC:
            text = text[:MAX_CHARS_PER_DOC] + "\n...[truncated]..."
        sections.append(f"--- Document: {f.name} ---\n{text}")

    return "\n\n".join(sections)


def extract_images(uploaded_files: List) -> List[Dict[str, Any]]:
    """Turn a list of image uploads into base64-encoded vision inputs.

    Returns a list of {"name", "mime_type", "data"} dicts, where "data" is
    a base64-encoded string ready to hand to any provider's vision input.
    Oversized images are skipped with a note rather than sent (or crashing).
    """
    images = []
    for f in uploaded_files:
        name_lower = f.name.lower()
        ext = next((e for e in IMAGE_EXTENSIONS if name_lower.endswith(e)), None)
        if ext is None:
            continue

        raw = f.getvalue()
        if len(raw) > MAX_IMAGE_BYTES:
            images.append(
                {
                    "name": f.name,
                    "mime_type": _IMAGE_MIME_TYPES[ext],
                    "data": None,
                    "error": f"Image is over the {MAX_IMAGE_BYTES // (1024 * 1024)}MB limit and was skipped.",
                }
            )
            continue

        images.append(
            {
                "name": f.name,
                "mime_type": _IMAGE_MIME_TYPES[ext],
                "data": base64.b64encode(raw).decode("ascii"),
            }
        )
    return images
