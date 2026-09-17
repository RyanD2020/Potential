"""Turns the generated plan (plus progress) into shareable files.

Two formats, for two different sharing needs:
- PDF: a clean, readable document - what you'd forward to a manager or
  attach to an email.
- Excel: a working spreadsheet - what you'd hand to someone who wants to
  filter/sort/track steps themselves in a tool they already use.
"""
from __future__ import annotations

import io
from typing import Any, Dict
from xml.sax.saxutils import escape

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)

BRAND_BLUE = "1E5FAE"


def _esc(text: Any) -> str:
    """Escape text for safe use inside a reportlab Paragraph's markup."""
    return escape(str(text or ""))


def _step_key(phase_idx: int, step_idx: int) -> str:
    return f"{phase_idx}-{step_idx}"


def build_plan_pdf(plan: Dict[str, Any], progress: Dict[str, bool]) -> bytes:
    """Render the plan as a PDF and return its bytes."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        title=plan.get("project_title", "Project Plan"),
    )

    styles = getSampleStyleSheet()
    title_style = styles["Title"]
    body_style = styles["BodyText"]
    phase_style = ParagraphStyle(
        "PhaseHeading", parent=styles["Heading2"], textColor=colors.HexColor(f"#{BRAND_BLUE}")
    )
    step_style = ParagraphStyle(
        "StepHeading", parent=styles["Heading3"], spaceBefore=10, spaceAfter=2
    )
    meta_style = ParagraphStyle(
        "Meta", parent=styles["BodyText"], fontSize=9, textColor=colors.HexColor("#5B6472")
    )

    phases = plan.get("phases", [])
    total_steps = sum(len(p.get("steps", [])) for p in phases)
    done_steps = sum(1 for v in progress.values() if v)

    elements = [
        Paragraph(_esc(plan.get("project_title", "Project Plan")), title_style),
        Spacer(1, 6),
    ]
    if plan.get("summary"):
        elements.append(Paragraph(_esc(plan["summary"]), body_style))
        elements.append(Spacer(1, 8))
    elements.append(Paragraph(f"Progress: {done_steps}/{total_steps} steps complete", meta_style))
    elements.append(Spacer(1, 14))

    for p_idx, phase in enumerate(phases):
        elements.append(
            Paragraph(f"Phase {p_idx + 1}: {_esc(phase.get('title', ''))}", phase_style)
        )
        if phase.get("goal"):
            elements.append(Paragraph(f"<i>Goal: {_esc(phase['goal'])}</i>", body_style))
        elements.append(Spacer(1, 4))

        for s_idx, step in enumerate(phase.get("steps", [])):
            done = progress.get(_step_key(p_idx, s_idx), False)
            status = "[x] Done" if done else "[ ] Not started"
            elements.append(
                Paragraph(f"{status} \u2014 {_esc(step.get('title', ''))}", step_style)
            )
            if step.get("description"):
                elements.append(Paragraph(_esc(step["description"]), body_style))
            if step.get("why_it_matters"):
                elements.append(
                    Paragraph(f"<i>Why it matters:</i> {_esc(step['why_it_matters'])}", meta_style)
                )
            questions = step.get("self_check_questions") or []
            if questions:
                items = [ListItem(Paragraph(_esc(q), meta_style)) for q in questions]
                elements.append(ListFlowable(items, bulletType="bullet", leftIndent=14))

        elements.append(Spacer(1, 10))

    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()


def build_plan_excel(plan: Dict[str, Any], progress: Dict[str, bool]) -> bytes:
    """Render the plan as an .xlsx workbook and return its bytes."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Plan"

    sheet.append([plan.get("project_title", "Project Plan")])
    sheet["A1"].font = Font(size=14, bold=True)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)

    if plan.get("summary"):
        sheet.append([plan["summary"]])
        sheet.merge_cells(
            start_row=sheet.max_row, start_column=1, end_row=sheet.max_row, end_column=6
        )
        sheet.cell(row=sheet.max_row, column=1).alignment = Alignment(wrap_text=True)

    sheet.append([])

    header = ["Phase", "Step", "Status", "Description", "Why It Matters", "Self-Check Questions"]
    sheet.append(header)
    header_row_idx = sheet.max_row
    header_fill = PatternFill(start_color=BRAND_BLUE, end_color=BRAND_BLUE, fill_type="solid")
    for col_idx in range(1, len(header) + 1):
        cell = sheet.cell(row=header_row_idx, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    for p_idx, phase in enumerate(plan.get("phases", [])):
        for s_idx, step in enumerate(phase.get("steps", [])):
            done = progress.get(_step_key(p_idx, s_idx), False)
            questions = "; ".join(step.get("self_check_questions") or [])
            sheet.append(
                [
                    f"Phase {p_idx + 1}: {phase.get('title', '')}",
                    step.get("title", ""),
                    "Done" if done else "Not started",
                    step.get("description", ""),
                    step.get("why_it_matters", ""),
                    questions,
                ]
            )

    widths = [26, 28, 12, 45, 35, 45]
    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = width

    for row in sheet.iter_rows(min_row=header_row_idx + 1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
